"""엑셀·워드·PPT 양식에 값을 채워 넣는 부분.

양식 파일 안에 `{{항목}}` 이라고 적어 두면 그 자리에 값이 들어갑니다.
서식(글꼴, 색, 표, 수식)은 그대로 두고 글자만 바꾸므로, 회사 양식을 그대로
쓸 수 있습니다.

지원하는 형식
  .xlsx / .xlsm : openpyxl  (수식이 들어 있는 칸은 건드리지 않습니다)
  .docx         : python-docx (본문, 표, 머리말/꼬리말)
  .pptx         : python-pptx (도형 안의 글, 표)
  .md / .txt    : 글자 치환

워드는 한 문장이 여러 조각(run)으로 쪼개져 저장되는 일이 잦아서,
`{{항목}}` 이 조각 사이에 걸치면 찾지 못합니다. 그래서 먼저 문단의 조각을
하나로 합친 뒤 바꿉니다. (가장 흔한 실패 원인입니다)
"""
from __future__ import annotations

import re
from pathlib import Path

PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")
SUPPORTED = {".xlsx", ".xlsm", ".docx", ".pptx", ".md", ".txt", ".csv"}


class Filler:
    """값을 채우면서 '채운 항목'과 '못 채운 항목'을 같이 기록합니다."""

    def __init__(self, values: dict) -> None:
        self.values = {str(k): str(v) for k, v in (values or {}).items()}
        self.filled: set[str] = set()
        self.missing: set[str] = set()

    def text(self, source: str) -> str:
        def replace(match: re.Match) -> str:
            key = match.group(1)
            value = self.values.get(key, "")
            if value.strip():
                self.filled.add(key)
                return value
            self.missing.add(key)
            return f"(미작성: {key})"

        return PLACEHOLDER.sub(replace, source)

    def has_placeholder(self, source: str) -> bool:
        return bool(source) and "{{" in source

    def report(self) -> dict:
        return {"filled": sorted(self.filled), "missing": sorted(self.missing)}


def find_fields(path: Path) -> list[str]:
    """양식 파일 안에 있는 {{항목}} 이름들을 모아 돌려줍니다."""
    suffix = path.suffix.lower()
    found: set[str] = set()

    if suffix in {".md", ".txt", ".csv"}:
        found.update(PLACEHOLDER.findall(path.read_text(encoding="utf-8")))
    elif suffix in {".xlsx", ".xlsm"}:
        import openpyxl

        book = openpyxl.load_workbook(path, keep_vba=suffix == ".xlsm")
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        found.update(PLACEHOLDER.findall(cell.value))
        book.close()
    elif suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        for paragraph in _docx_paragraphs(document):
            found.update(PLACEHOLDER.findall(paragraph.text))
    elif suffix == ".pptx":
        import pptx

        presentation = pptx.Presentation(str(path))
        for paragraph in _pptx_paragraphs(presentation):
            found.update(PLACEHOLDER.findall(paragraph.text))
    return sorted(found)


def fill_file(source: Path, target: Path, values: dict) -> dict:
    """양식 파일을 채워 새 파일로 저장합니다. 원본은 그대로 둡니다."""
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(
            f"{suffix} 형식은 아직 자동 채우기를 못 합니다. 지원: {', '.join(sorted(SUPPORTED))}"
        )

    target.parent.mkdir(parents=True, exist_ok=True)
    filler = Filler(values)

    if suffix in {".md", ".txt", ".csv"}:
        target.write_text(filler.text(source.read_text(encoding="utf-8")), encoding="utf-8")
    elif suffix in {".xlsx", ".xlsm"}:
        _fill_xlsx(source, target, filler)
    elif suffix == ".docx":
        _fill_docx(source, target, filler)
    else:
        _fill_pptx(source, target, filler)

    return {"path": str(target), "filename": target.name, **filler.report()}


# ── 엑셀 ────────────────────────────────────────────────────────────────
def _fill_xlsx(source: Path, target: Path, filler: Filler) -> None:
    import openpyxl
    from openpyxl.cell.cell import MergedCell

    book = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                # 합쳐진 칸은 왼쪽 위 칸에만 쓸 수 있습니다.
                if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                    continue
                # 수식은 건드리지 않습니다. 계산 결과가 깨지면 안 되니까요.
                if cell.value.startswith("=") or not filler.has_placeholder(cell.value):
                    continue
                cell.value = filler.text(cell.value)
    book.save(target)
    book.close()


# ── 워드 ────────────────────────────────────────────────────────────────
def _docx_paragraphs(document) -> list:
    """본문, 표, 머리말/꼬리말에 있는 문단을 전부 모읍니다."""
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(_docx_table_paragraphs(table))
    for section in document.sections:
        for part in (section.header, section.footer):
            paragraphs.extend(part.paragraphs)
            for table in part.tables:
                paragraphs.extend(_docx_table_paragraphs(table))
    return paragraphs


def _docx_table_paragraphs(table) -> list:
    paragraphs = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
            for inner in cell.tables:
                paragraphs.extend(_docx_table_paragraphs(inner))
    return paragraphs


def _fill_docx(source: Path, target: Path, filler: Filler) -> None:
    import docx

    document = docx.Document(str(source))
    for paragraph in _docx_paragraphs(document):
        if not filler.has_placeholder(paragraph.text):
            continue
        replaced = filler.text(paragraph.text)
        runs = paragraph.runs
        if not runs:
            continue
        # 첫 조각에 전체 문장을 넣고 나머지는 비웁니다. 첫 조각의 서식이 유지됩니다.
        runs[0].text = replaced
        for run in runs[1:]:
            run.text = ""
    document.save(str(target))


# ── PPT ─────────────────────────────────────────────────────────────────
def _pptx_paragraphs(presentation) -> list:
    paragraphs = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            paragraphs.extend(_pptx_shape_paragraphs(shape))
    return paragraphs


def _pptx_shape_paragraphs(shape) -> list:
    paragraphs = []
    if shape.shape_type is not None and getattr(shape, "shapes", None) is not None:
        for inner in shape.shapes:  # 그룹으로 묶인 도형
            paragraphs.extend(_pptx_shape_paragraphs(inner))
    if getattr(shape, "has_text_frame", False):
        paragraphs.extend(shape.text_frame.paragraphs)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.text_frame.paragraphs)
    return paragraphs


def _fill_pptx(source: Path, target: Path, filler: Filler) -> None:
    import pptx

    presentation = pptx.Presentation(str(source))
    for paragraph in _pptx_paragraphs(presentation):
        if not filler.has_placeholder(paragraph.text):
            continue
        replaced = filler.text(paragraph.text)
        runs = paragraph.runs
        if not runs:
            continue
        runs[0].text = replaced
        for run in runs[1:]:
            run.text = ""
    presentation.save(str(target))

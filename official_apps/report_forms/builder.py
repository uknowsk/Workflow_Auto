"""고른 옵션대로 보고서 파일을 만들어 내는 부분.

`common/office.py` 는 **이미 있는 회사 양식 파일에 값을 채우는** 도구이고,
여기는 양식 파일이 아예 없을 때 **처음부터 만들어 내는** 쪽입니다.
나중에 진짜 회사 양식 파일이 생기면 office.fill_file 로 갈아타면 됩니다.

만들 수 있는 형식: 워드(.docx), 엑셀(.xlsx), PPT(.pptx), HTML, 마크다운(.md)
HTML 은 사내 폐쇄망을 생각해 바깥에서 아무것도 받아오지 않습니다
(글꼴은 기기에 있는 것만, 그림·스크립트 없음).
"""
from __future__ import annotations

from pathlib import Path

# 사용자가 뭐라고 부르든 알아듣게 합니다.
FORMATS: dict[str, str] = {
    "word": "docx", "워드": "docx", "docx": "docx", "doc": "docx", "한글": "docx",
    "excel": "xlsx", "엑셀": "xlsx", "xlsx": "xlsx",
    "ppt": "pptx", "파워포인트": "pptx", "pptx": "pptx", "슬라이드": "pptx",
    "html": "html", "웹": "html", "웹페이지": "html",
    "md": "md", "markdown": "md", "마크다운": "md", "텍스트": "md", "text": "md",
}
# 글로 돌려줄 수 있는 형식(파일을 내려받지 않고 바로 붙여넣을 수 있는 것)
TEXT_FORMATS = {"html", "md"}


def resolve_format(name: str) -> str:
    return FORMATS.get((name or "").strip().lower(), "")


def _body_or_guide(section: dict) -> tuple[str, bool]:
    """내용이 있으면 내용을, 없으면 작성 안내를 돌려줍니다. (글, 안내인가)"""
    body = (section.get("body") or "").strip()
    if body:
        return body, False
    return section.get("guide", ""), True


def _lines(text: str) -> list[str]:
    return [line.strip() for line in (text or "").splitlines() if line.strip()]


# ── 마크다운 ────────────────────────────────────────────────────────────
def build_md(doc: dict) -> str:
    out = [f"# {doc['title']}", ""]
    for label, value in doc["header"]:
        out.append(f"- {label}: {value or '(미작성)'}")
    out.append("")
    for section in doc["sections"]:
        out.append(f"## {section['title']}")
        text, is_guide = _body_or_guide(section)
        out.append(f"> 작성 안내: {text}" if is_guide else text)
        out.append("")
    return "\n".join(out).rstrip() + "\n"


# ── HTML ───────────────────────────────────────────────────────────────
_CSS = """
:root { --line:#d9dce1; --ink:#1b1d21; --dim:#6b7280; }
* { box-sizing:border-box; }
body { margin:0; padding:32px 24px; color:var(--ink); background:#fff;
  font-family:-apple-system,BlinkMacSystemFont,"Malgun Gothic","맑은 고딕",
  "Apple SD Gothic Neo","Noto Sans KR",sans-serif; line-height:1.7;
  word-break:keep-all; overflow-wrap:break-word; }
.sheet { max-width:820px; margin:0 auto; }
h1 { font-size:24px; margin:0 0 20px; padding-bottom:14px; border-bottom:2px solid var(--ink); }
table.meta { width:100%; border-collapse:collapse; margin-bottom:28px; font-size:14px; }
table.meta th, table.meta td { border:1px solid var(--line); padding:7px 10px; text-align:left; }
table.meta th { width:110px; background:#f6f7f9; font-weight:600; color:var(--dim); }
h2 { font-size:17px; margin:26px 0 8px; padding-left:10px; border-left:4px solid var(--ink); }
p { margin:0 0 10px; white-space:pre-wrap; }
p.guide { color:var(--dim); font-size:14px; background:#f6f7f9; border:1px dashed var(--line);
  border-radius:6px; padding:10px 12px; }
footer { margin-top:32px; padding-top:12px; border-top:1px solid var(--line);
  color:var(--dim); font-size:12px; }
@media print { body { padding:0; } footer { display:none; } h2 { break-after:avoid; } }
@page { size:A4; margin:18mm; }
"""


def _escape(text: str) -> str:
    return (
        str(text)
        .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    )


def build_html(doc: dict) -> str:
    rows = "\n".join(
        f"    <tr><th>{_escape(label)}</th><td>{_escape(value or '')}</td></tr>"
        for label, value in doc["header"]
    )
    blocks = []
    for section in doc["sections"]:
        text, is_guide = _body_or_guide(section)
        css = ' class="guide"' if is_guide else ""
        prefix = "작성 안내 · " if is_guide else ""
        blocks.append(
            f"  <h2>{_escape(section['title'])}</h2>\n"
            f"  <p{css}>{prefix}{_escape(text)}</p>"
        )
    return (
        "<!doctype html>\n<html lang=\"ko\">\n<head>\n<meta charset=\"utf-8\">\n"
        f"<title>{_escape(doc['title'])}</title>\n<style>{_CSS}</style>\n</head>\n"
        "<body>\n<div class=\"sheet\">\n"
        f"  <h1>{_escape(doc['title'])}</h1>\n  <table class=\"meta\">\n{rows}\n  </table>\n"
        + "\n".join(blocks)
        + f"\n  <footer>{_escape(doc['report_type'])} · {_escape(doc['length'])} 분량 양식</footer>\n"
        "</div>\n</body>\n</html>\n"
    )


# ── 워드 ───────────────────────────────────────────────────────────────
def build_docx(doc: dict, target: Path) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.add_heading(doc["title"], level=0)

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in doc["header"]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value or "")
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    document.add_paragraph("")

    for section in doc["sections"]:
        document.add_heading(section["title"], level=1)
        text, is_guide = _body_or_guide(section)
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"작성 안내 · {text}" if is_guide else text)
        if is_guide:
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        document.add_paragraph("")

    document.save(target)
    return target


# ── 엑셀 ───────────────────────────────────────────────────────────────
def build_xlsx(doc: dict, target: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = Workbook()
    sheet = book.active
    sheet.title = "보고서"
    thin = Side(style="thin", color="D9DCE1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="F6F7F9")
    wrap = Alignment(vertical="top", wrap_text=True)

    sheet["A1"] = doc["title"]
    sheet["A1"].font = Font(size=16, bold=True)
    sheet.merge_cells("A1:C1")

    row = 3
    for label, value in doc["header"]:
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=row, column=1).fill = head_fill
        sheet.cell(row=row, column=2, value=str(value or ""))
        for column in (1, 2, 3):
            sheet.cell(row=row, column=column).border = border
            sheet.cell(row=row, column=column).alignment = wrap
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        row += 1

    row += 1
    for column, title in enumerate(("항목", "내용", "작성 안내"), start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.border = border
    row += 1

    for section in doc["sections"]:
        sheet.cell(row=row, column=1, value=section["title"]).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=(section.get("body") or "").strip())
        sheet.cell(row=row, column=3, value=section.get("guide", "")).font = Font(
            size=9, color="6B7280", italic=True
        )
        for column in (1, 2, 3):
            sheet.cell(row=row, column=column).border = border
            sheet.cell(row=row, column=column).alignment = wrap
        sheet.row_dimensions[row].height = 54
        row += 1

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 62
    sheet.column_dimensions["C"].width = 40
    book.save(target)
    return target


# ── PPT ────────────────────────────────────────────────────────────────
def build_pptx(doc: dict, target: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Pt

    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = doc["title"]
    cover.placeholders[1].text = " · ".join(
        f"{label} {value}" for label, value in doc["header"] if value
    ) or f"{doc['report_type']} ({doc['length']} 분량)"

    for section in doc["sections"]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = section["title"]
        frame = slide.placeholders[1].text_frame
        frame.word_wrap = True
        text, is_guide = _body_or_guide(section)
        lines = _lines(text) or [""]
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = f"작성 안내 · {line}" if is_guide else line
            paragraph.font.size = Pt(16 if not is_guide else 14)
            if is_guide:
                paragraph.font.italic = True

    deck.save(target)
    return target


BUILDERS = {"docx": build_docx, "xlsx": build_xlsx, "pptx": build_pptx}


def build(doc: dict, file_format: str, target_dir: Path, stem: str) -> dict:
    """형식에 맞게 만들어, 글이면 글을 / 파일이면 경로를 돌려줍니다."""
    if file_format == "md":
        return {"format": "md", "document": build_md(doc)}
    if file_format == "html":
        text = build_html(doc)
        target_dir.mkdir(parents=True, exist_ok=True)
        path = target_dir / f"{stem}.html"
        path.write_text(text, encoding="utf-8")
        # HTML 은 붙여넣기도 하고 파일로도 받으므로 둘 다 돌려줍니다.
        return {"format": "html", "document": text, "path": path}
    if file_format not in BUILDERS:
        return {"error": f"만들 수 없는 형식입니다: {file_format}"}
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / f"{stem}.{file_format}"
    BUILDERS[file_format](doc, path)
    return {"format": file_format, "path": path}
